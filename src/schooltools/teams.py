"""Lettura dell'elenco di origine per l'importazione utenti di Microsoft 365.

Il file di partenza e' un xlsx con le colonne `Nome`, `Cognome` e `Classe`,
cioe' quello che si esporta di solito dalla segreteria. Non se ne puo' dare per
scontato il tracciato: l'intestazione puo' non essere sulla prima riga, le
colonne possono stare in un ordine qualunque e in mezzo ai dati ci sono righe
vuote da saltare. Tutto cio' che manca diventa un `ValueError` con un messaggio
leggibile, mai un `IndexError` o un `KeyError` che risale all'utente.
"""

from typing import TYPE_CHECKING
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from schooltools.students import Studente

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path

    from openpyxl.workbook.workbook import Workbook

COLONNE_ORIGINE = ("Nome", "Cognome", "Classe")


def leggi_origine(percorso: Path) -> list[Studente]:
    """Legge gli studenti da un xlsx con le colonne `Nome`, `Cognome`, `Classe`.

    Args:
        percorso: File xlsx da leggere.

    Returns:
        Gli studenti nell'ordine del file, senza le righe vuote.

    Raises:
        ValueError: Se il file non e' un xlsx leggibile, se non ha fogli, se
            manca una delle colonne o se una riga piena non ha ne' nome ne'
            cognome.
        OSError: Se il file non esiste o non e' leggibile.
    """
    cartella = _apri(percorso)
    try:
        if not cartella.worksheets:
            raise ValueError(f"{percorso} non contiene nessun foglio")
        foglio = cartella.worksheets[0]
        righe = [[_testo(cella) for cella in riga] for riga in foglio.iter_rows(values_only=True)]
    finally:
        cartella.close()
    return studenti_da_righe(righe)


def studenti_da_righe(righe: Sequence[Sequence[str]]) -> list[Studente]:
    """Ricava gli studenti dalle celle del foglio di origine, gia' ridotte a testo.

    Args:
        righe: Righe del foglio, una lista di celle ciascuna.

    Returns:
        Gli studenti nell'ordine del file, senza le righe vuote.

    Raises:
        ValueError: Se manca l'intestazione o una delle colonne, oppure se una
            riga piena non ha ne' nome ne' cognome.
    """
    inizio, colonne = _intestazione(righe)
    studenti = []
    for numero, riga in enumerate(righe[inizio + 1 :], inizio + 2):
        if not any(riga):
            continue
        studente = Studente(
            nome=_campo(riga, colonne["nome"]),
            cognome=_campo(riga, colonne["cognome"]),
            classe=_campo(riga, colonne["classe"]),
        )
        if not studente.nome and not studente.cognome:
            raise ValueError(f"La riga {numero} non ha ne' nome ne' cognome")
        studenti.append(studente)
    return studenti


def _apri(percorso: Path) -> Workbook:
    """Apre il file xlsx, traducendo in ValueError i modi in cui non lo e'."""
    try:
        return load_workbook(percorso, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, ValueError) as errore:
        raise ValueError(f"{percorso} non e' un file xlsx leggibile") from errore


def _intestazione(righe: Sequence[Sequence[str]]) -> tuple[int, dict[str, int]]:
    """La riga d'intestazione e la colonna di ognuno dei nomi che servono."""
    cercate = {nome.casefold() for nome in COLONNE_ORIGINE}
    parziale: tuple[int, dict[str, int]] | None = None
    for numero, riga in enumerate(righe):
        trovate: dict[str, int] = {}
        for i, cella in enumerate(riga):
            chiave = cella.casefold()
            if chiave in cercate:
                trovate.setdefault(chiave, i)
        if len(trovate) == len(cercate):
            return numero, trovate
        if trovate and parziale is None:
            parziale = (numero, trovate)

    mancanti = [c for c in COLONNE_ORIGINE if parziale is None or c.casefold() not in parziale[1]]
    raise ValueError(f"Nel file di origine mancano le colonne: {', '.join(mancanti)}")


def _campo(riga: Sequence[str], i: int) -> str:
    """Il valore della colonna `i`, vuoto se la riga finisce prima."""
    return riga[i] if i < len(riga) else ""


def _testo(valore: object) -> str:
    """Il contenuto di una cella come testo ripulito, vuoto se la cella e' vuota."""
    return "" if valore is None else str(valore).strip()
