"""Test della lettura dell'elenco di origine per l'importazione utenti."""

from typing import TYPE_CHECKING

import pytest
from openpyxl import Workbook

from schooltools.students import Studente
from schooltools.teams import leggi_origine, studenti_da_righe

if TYPE_CHECKING:
    from collections.abc import Sequence
    from pathlib import Path


def _origine(percorso: Path, righe: Sequence[Sequence[object]]) -> Path:
    """Scrive un file di origine con le righe date, cosi' come sono."""
    cartella = Workbook()
    foglio = cartella.active
    assert foglio is not None
    for riga in righe:
        foglio.append(list(riga))
    cartella.save(percorso)
    return percorso


def test_leggi_origine(tmp_path: Path) -> None:
    """Le tre colonne diventano uno `Studente` per riga, nell'ordine del file."""
    percorso = _origine(
        tmp_path / "origine.xlsx",
        [
            ["Nome", "Cognome", "Classe"],
            ["Giulio", "Rossi", "1E LSA"],
            ["Maria Luisa", "D'Agostino", "5E LSA"],
        ],
    )

    assert leggi_origine(percorso) == [
        Studente(nome="Giulio", cognome="Rossi", classe="1E LSA"),
        Studente(nome="Maria Luisa", cognome="D'Agostino", classe="5E LSA"),
    ]


def test_leggi_origine_salta_le_righe_vuote(tmp_path: Path) -> None:
    """Le righe senza niente dentro non diventano studenti."""
    percorso = _origine(
        tmp_path / "origine.xlsx",
        [
            ["Nome", "Cognome", "Classe"],
            ["Giulio", "Rossi", "1E"],
            [None, None, None],
            ["", "  ", ""],
            ["Anna", "Verdi", "2B"],
            [None, None, None],
        ],
    )

    assert [s.cognome for s in leggi_origine(percorso)] == ["Rossi", "Verdi"]


def test_leggi_origine_con_celle_non_testuali(tmp_path: Path) -> None:
    """Una classe scritta come numero arriva lo stesso come testo."""
    percorso = _origine(
        tmp_path / "origine.xlsx",
        [["Nome", "Cognome", "Classe"], ["Giulio", "Rossi", 1]],
    )

    assert leggi_origine(percorso) == [Studente(nome="Giulio", cognome="Rossi", classe="1")]


def test_leggi_origine_file_non_xlsx(tmp_path: Path) -> None:
    """Un file che non e' un xlsx diventa un errore leggibile, non un traceback."""
    percorso = tmp_path / "origine.xlsx"
    percorso.write_text("Nome,Cognome,Classe\n", encoding="utf-8")

    with pytest.raises(ValueError, match="non e' un file xlsx leggibile"):
        leggi_origine(percorso)


def test_leggi_origine_file_assente(tmp_path: Path) -> None:
    """Un file che non esiste resta un errore di sistema, gestito dalla CLI."""
    with pytest.raises(OSError, match="assente"):
        leggi_origine(tmp_path / "assente.xlsx")


def test_leggi_origine_senza_fogli(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Una cartella senza nessun foglio e' un errore, non un IndexError."""
    vuota = Workbook()
    foglio = vuota.active
    assert foglio is not None
    vuota.remove(foglio)

    def apri(*_: object, **__: object) -> Workbook:
        return vuota

    monkeypatch.setattr("schooltools.teams.load_workbook", apri)

    with pytest.raises(ValueError, match="nessun foglio"):
        leggi_origine(tmp_path / "origine.xlsx")


def test_studenti_da_righe_colonne_in_ordine_qualunque() -> None:
    """Le colonne si cercano per nome, non per posizione."""
    righe = [["Classe", "Cognome", "Nome"], ["1E", "Rossi", "Giulio"]]

    assert studenti_da_righe(righe) == [Studente(nome="Giulio", cognome="Rossi", classe="1E")]


def test_studenti_da_righe_intestazione_non_in_cima() -> None:
    """Titoli e righe vuote prima dell'intestazione non danno fastidio."""
    righe = [
        ["Elenco studenti"],
        ["", "", ""],
        ["Nome", "Cognome", "Classe"],
        ["Anna", "Verdi", ""],
    ]

    assert studenti_da_righe(righe) == [Studente(nome="Anna", cognome="Verdi", classe="")]


def test_studenti_da_righe_intestazione_con_maiuscole_diverse() -> None:
    """L'intestazione si riconosce a prescindere dalle maiuscole."""
    righe = [["NOME", "cognome", "Classe"], ["Anna", "Verdi", "2B"]]

    assert studenti_da_righe(righe)[0].nome == "Anna"


def test_studenti_da_righe_riga_piu_corta_dell_intestazione() -> None:
    """Una riga che finisce prima delle colonne da' campi vuoti, non un IndexError."""
    righe = [["Nome", "Cognome", "Classe"], ["Anna", "Verdi"]]

    assert studenti_da_righe(righe) == [Studente(nome="Anna", cognome="Verdi", classe="")]


def test_studenti_da_righe_colonna_mancante() -> None:
    """Se manca una colonna il messaggio dice quale."""
    righe = [["Nome", "Cognome"], ["Anna", "Verdi"]]

    with pytest.raises(ValueError, match="mancano le colonne: Classe"):
        studenti_da_righe(righe)


def test_studenti_da_righe_senza_intestazione() -> None:
    """Senza nessuna intestazione riconoscibile si elencano tutte le colonne attese."""
    righe = [["Anna", "Verdi", "2B"]]

    with pytest.raises(ValueError, match="mancano le colonne: Nome, Cognome, Classe"):
        studenti_da_righe(righe)


def test_studenti_da_righe_riga_senza_nome_ne_cognome() -> None:
    """Una riga piena ma senza nome ne' cognome e' segnalata con il suo numero."""
    righe = [["Nome", "Cognome", "Classe"], ["Anna", "Verdi", "2B"], ["", "", "2B"]]

    with pytest.raises(ValueError, match="riga 3"):
        studenti_da_righe(righe)


def test_studenti_da_righe_senza_studenti() -> None:
    """Un file con la sola intestazione non da' nessuno studente."""
    assert studenti_da_righe([["Nome", "Cognome", "Classe"]]) == []
