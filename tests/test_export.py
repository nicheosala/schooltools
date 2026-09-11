"""Test della scrittura degli elenchi di studenti in xlsx."""

from typing import TYPE_CHECKING

import pytest
from openpyxl import load_workbook

from schooltools.export import INTESTAZIONI, nome_foglio, scrivi_utenti_xlsx, scrivi_xlsx
from schooltools.students import CAMPI_UTENTI, Studente

if TYPE_CHECKING:
    from pathlib import Path


def test_nome_foglio_corto_resta_uguale() -> None:
    """Un nome che rientra nei limiti di xlsx non viene toccato."""
    assert nome_foglio("1A") == "1A"


def test_nome_foglio_abbrevia_i_nomi_lunghi() -> None:
    """Oltre i 31 caratteri si tiene la sola sigla della classe."""
    assert nome_foglio("1E LSA LICEO SCIENTIFICO OPZIONE SCIENZE APPLICATE") == "1E LSA"


def test_nome_foglio_tronca_quando_non_puo_abbreviare() -> None:
    """Senza una sigla da isolare il nome viene tagliato a 31 caratteri."""
    nome = nome_foglio("CORSO SERALE DI INFORMATICA PER ADULTI LAVORATORI")

    assert len(nome) == 31
    assert nome.startswith("CORSO SERALE")


def test_nome_foglio_lascia_stare_i_nomi_gia_corti() -> None:
    """Un nome che non supera il limite non viene abbreviato, anche se ha piu' parole."""
    assert nome_foglio("3A CORSO SERALE") == "3A CORSO SERALE"


def test_nome_foglio_toglie_i_caratteri_vietati() -> None:
    """I caratteri che xlsx non ammette in un titolo diventano spazi."""
    assert nome_foglio("3A/B: [prova]?*") == "3A B prova"


def test_nome_foglio_evita_i_duplicati() -> None:
    """Un titolo gia' assegnato riceve un numero progressivo."""
    assert nome_foglio("1A", usati=["1A"]) == "1A (2)"
    assert nome_foglio("1A", usati=["1A", "1A (2)"]) == "1A (3)"


def test_nome_foglio_vuoto() -> None:
    """Un nome fatto di soli spazi diventa "Classe"."""
    assert nome_foglio("   ") == "Classe"


def test_scrivi_xlsx_un_foglio_per_classe(tmp_path: Path) -> None:
    """Ogni classe ha il suo foglio, con l'intestazione e una riga per studente."""
    percorso = scrivi_xlsx(
        tmp_path / "studenti.xlsx",
        {"1A": ["Rossi Giulio", "D'Agostino Maria Luisa"], "2B": ["Verdi Anna"]},
    )

    cartella = load_workbook(percorso)
    assert cartella.sheetnames == ["1A", "2B"]

    foglio = cartella["1A"]
    assert [c.value for c in foglio[1]] == list(INTESTAZIONI)
    assert [[c.value for c in riga] for riga in foglio.iter_rows(min_row=2)] == [
        ["Rossi", "Giulio"],
        ["D'Agostino", "Maria Luisa"],
    ]
    assert cartella["2B"].max_row == 2


def test_scrivi_xlsx_classi_con_nome_lungo_e_uguale(tmp_path: Path) -> None:
    """Due classi che si accorciano allo stesso titolo restano su fogli distinti."""
    lungo = "LICEO SCIENTIFICO OPZIONE SCIENZE APPLICATE"
    percorso = scrivi_xlsx(tmp_path / "s.xlsx", {f"{lungo} A": [], f"{lungo} B": []})

    nomi = load_workbook(percorso).sheetnames
    assert len(nomi) == 2
    assert all(len(n) <= 31 for n in nomi)


def test_scrivi_xlsx_crea_le_cartelle(tmp_path: Path) -> None:
    """Le cartelle mancanti nel percorso di destinazione vengono create."""
    percorso = scrivi_xlsx(tmp_path / "a" / "b" / "s.xlsx", {"1A": ["Rossi Giulio"]})

    assert percorso.is_file()


def test_scrivi_xlsx_senza_classi(tmp_path: Path) -> None:
    """Senza nessuna classe non si scrive un file vuoto: si segnala l'errore."""
    with pytest.raises(ValueError, match="Nessuna classe"):
        scrivi_xlsx(tmp_path / "s.xlsx", {})


def test_scrivi_utenti_xlsx(tmp_path: Path) -> None:
    """Il foglio ha le colonne di Microsoft 365 e una riga per studente."""
    percorso = scrivi_utenti_xlsx(
        tmp_path / "utenti.xlsx",
        [
            Studente(nome="Giulio", cognome="Rossi", classe="1E"),
            Studente(nome="Maria Luisa", cognome="D'Agostino", classe="5E"),
        ],
    )

    foglio = load_workbook(percorso).active
    assert foglio is not None
    assert [c.value for c in foglio[1]] == CAMPI_UTENTI
    assert [c.value for c in foglio[2]][:6] == [
        "giulio.rossi@istitutobachelet.edu.it",
        "Giulio",
        "Rossi",
        "Giulio Rossi",
        "Studente",
        "1E",
    ]
    assert foglio.max_row == 3


def test_scrivi_utenti_xlsx_crea_le_cartelle(tmp_path: Path) -> None:
    """Le cartelle mancanti nel percorso di destinazione vengono create."""
    percorso = scrivi_utenti_xlsx(tmp_path / "a" / "b" / "u.xlsx", [Studente("Anna", "Verdi")])

    assert percorso.is_file()


def test_scrivi_utenti_xlsx_senza_studenti(tmp_path: Path) -> None:
    """Senza nessuno studente non si scrive un file vuoto: si segnala l'errore."""
    with pytest.raises(ValueError, match="Nessuno studente"):
        scrivi_utenti_xlsx(tmp_path / "u.xlsx", [])
