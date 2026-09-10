from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

from schooltools.cli import _chiedi_voce, _scegli_classi, main
from schooltools.parsing import Classe

CLASSI = [Classe(id="1", nome="1E LSA"), Classe(id="2", nome="5E LSA")]


@pytest.fixture
def pagina(tmp_path: Path, registro_html: str) -> Path:
    percorso = tmp_path / "classe.html"
    percorso.write_text(registro_html, encoding="utf-8")
    return percorso


def _risposte(monkeypatch: pytest.MonkeyPatch, *valori: str) -> None:
    coda: Iterator[str] = iter(valori)
    monkeypatch.setattr("builtins.input", lambda _="": next(coda))


def test_scegli_classi_tutte() -> None:
    assert _scegli_classi(CLASSI, "TUTTE") == CLASSI


def test_scegli_classi_per_nome_parziale() -> None:
    assert _scegli_classi(CLASSI, "5e") == [CLASSI[1]]


def test_scegli_classi_senza_corrispondenze() -> None:
    with pytest.raises(ValueError, match="Nessuna classe"):
        _scegli_classi(CLASSI, "3Z")


def test_scegli_classi_interattivo(monkeypatch: pytest.MonkeyPatch) -> None:
    _risposte(monkeypatch, "2")

    assert _scegli_classi(CLASSI, None) == [CLASSI[0]]


def test_scegli_classi_interattivo_tutte(monkeypatch: pytest.MonkeyPatch) -> None:
    _risposte(monkeypatch, "")

    assert _scegli_classi(CLASSI, None) == CLASSI


def test_chiedi_voce_insiste_finche_la_risposta_non_e_valida(
    monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    _risposte(monkeypatch, "9", "zero", "2")

    assert _chiedi_voce("Quale?", ["a", "b"]) == 1
    assert "Inserisci un numero da 1 a 2." in capsys.readouterr().out


def test_main_da_file_mostra_i_nomi(
    pagina: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    codice = main(["--file", str(pagina), "--schermo"])

    uscita = capsys.readouterr().out
    assert codice == 0
    assert "1E LSA" in uscita
    assert "1. Rossi Giulio" in uscita
    assert "3. Verdi Anna" in uscita


def test_main_da_file_scrive_xlsx(
    pagina: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    destinazione = tmp_path / "studenti.xlsx"

    codice = main(["--file", str(pagina), "--xlsx", str(destinazione)])

    assert codice == 0
    assert "Scritti 3 studenti" in capsys.readouterr().out
    foglio = load_workbook(destinazione).active
    assert foglio is not None
    assert foglio.max_row == 4


def test_main_chiede_cosa_fare(
    pagina: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destinazione = tmp_path / "scelto.xlsx"
    _risposte(monkeypatch, "2", str(destinazione))

    assert main(["--file", str(pagina)]) == 0
    assert destinazione.is_file()


def test_main_senza_credenziali(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    monkeypatch.delenv("SPAGGIARI_UTENTE", raising=False)
    monkeypatch.delenv("SPAGGIARI_PASSWORD", raising=False)

    codice = main(["--env", str(tmp_path / "assente.env"), "--schermo"])

    assert codice == 1
    assert "Credenziali non trovate" in capsys.readouterr().out


def test_main_annullato(pagina: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def interrompi(_: str = "") -> str:
        raise KeyboardInterrupt

    monkeypatch.setattr("builtins.input", interrompi)

    assert main(["--file", str(pagina)]) == 1


def test_main_con_file_inesistente(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    codice = main(["--file", str(tmp_path / "assente.html"), "--schermo"])

    assert codice == 1
    assert "Errore:" in capsys.readouterr().out
